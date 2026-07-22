#!/usr/bin/env python3
import openpyxl
import requests
import re
import time

EXCEL_PATH   = '/home/wvieira/projetos/GOSP/produtos-vendidos-atualizado.xlsx'
VOUCHER_BASE = 'https://estacionamento-beta-rosy.vercel.app/img/{ident}.jpg'
API_URL      = 'https://atendimento.gosp.org.br/api/integrations/whatsapp/congresso/confirmacao'
API_TOKEN    = 'BYWdkDPgE0L3JvEm/6S0CZo7Ydp0wYyNlaSHM/SKI+u24g0HUJEatwpd/EZL2gPR'

TEST_MODE = False
TEST_CIM  = 282677   # CIM que recebe todas as mensagens no modo teste

DELAY_SECONDS = 2


def sanitize(v):
    return re.sub(r'\D', '', str(v))


def enviar(cim: int, url: str) -> dict:
    resp = requests.post(
        API_URL,
        headers={
            'Authorization': f'Bearer {API_TOKEN}',
            'Content-Type': 'application/json',
        },
        json={'cim': cim, 'url': url},
        timeout=15,
    )
    return resp.json()


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    entries = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[col['Nome']]).strip()
        cim  = sanitize(str(row[col['CIM']]))
        cpf  = sanitize(str(row[col['CPF']]))

        if name.lower() in seen:
            continue

        # CIM válido: até 7 dígitos
        if cim and len(cim) <= 7:
            ident     = cim
            cim_int   = int(cim)
            usar_cpf  = False
        elif cpf:
            ident     = cpf
            cim_int   = None
            usar_cpf  = True
        else:
            continue

        seen.add(name.lower())
        entries.append((name, cim_int, ident, usar_cpf))

    total = len(entries)
    print(f'{total} destinatários encontrados.')

    if TEST_MODE:
        print(f'\n*** MODO TESTE — todas as mensagens vão para CIM {TEST_CIM} ***\n')
        # Envia apenas o primeiro como exemplo
        name, cim_int, ident, usar_cpf = entries[0]
        url = VOUCHER_BASE.format(ident=ident)
        resp = enviar(TEST_CIM, url)
        print(f'[TESTE] {name} | url={url}')
        print(f'Resposta: {resp}')
        print('\nVerifique o WhatsApp. Quando aprovado, defina TEST_MODE=False e rode novamente.')
        return

    sent = 0
    nao_enviados = []

    for i, (name, cim_int, ident, usar_cpf) in enumerate(entries):
        if cim_int is None:
            print(f'  [SKIP] {name} — sem CIM')
            nao_enviados.append((name, '-', 'Sem CIM'))
            continue

        url = VOUCHER_BASE.format(ident=ident)
        try:
            resp = enviar(cim_int, url)
            if resp.get('ok'):
                print(f'[{i+1}/{total}] ✓ {name} | CIM={cim_int}')
                sent += 1
            else:
                print(f'[{i+1}/{total}] ✗ {name} | CIM={cim_int} | {resp}')
                nao_enviados.append((name, cim_int, str(resp)))
        except Exception as e:
            print(f'[{i+1}/{total}] [ERRO] {name} | CIM={cim_int} | {e}')
            nao_enviados.append((name, cim_int, str(e)))

        time.sleep(DELAY_SECONDS)

    # Salva lista de não enviados
    log_path = '/home/wvieira/projetos/GOSP/whatsapp_nao_enviados.txt'
    with open(log_path, 'w') as f:
        f.write(f'Total não enviados: {len(nao_enviados)}\n\n')
        f.write(f'{"Nome":<50} {"CIM":<12} Motivo\n')
        f.write('-' * 90 + '\n')
        for name, cim, motivo in nao_enviados:
            f.write(f'{name:<50} {str(cim):<12} {motivo}\n')

    print(f'\n✓ {sent} enviados | ✗ {len(nao_enviados)} não enviados')
    print(f'Lista salva em: {log_path}')


if __name__ == '__main__':
    main()
