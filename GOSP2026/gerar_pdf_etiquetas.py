#!/usr/bin/env python3
import openpyxl
import qrcode
import re
import io
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

EXCEL_PATH = '/home/wvieira/projetos/GOSP/produtos-vendidos-atualizado.xlsx'
OUTPUT_PDF = '/home/wvieira/projetos/GOSP/etiquetas_qrcode.pdf'
FONT_PATH  = '/usr/share/fonts/noto/NotoSerif-Bold.ttf'
FONT_SANS  = '/usr/share/fonts/noto/NotoSans-Regular.ttf'

# Layout
COLS        = 3
MARGIN      = 8 * mm
PAGE_W, PAGE_H = A4                      # 595.27 x 841.89 pt
CELL_W      = (PAGE_W - 2 * MARGIN) / COLS
CELL_H      = 55 * mm                    # altura de cada etiqueta
QR_SIZE     = 38 * mm
ROWS_PER_PAGE = int((PAGE_H - 2 * MARGIN) / CELL_H)


def sanitize_cim(cim):
    return re.sub(r'\D', '', cim)


def make_qr_image(data, size_pt):
    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M,
                       box_size=10, border=1)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color='black', back_color='white').convert('RGB')
    img = img.resize((int(size_pt * 3), int(size_pt * 3)), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf


def fit_font_size(c, text, max_width, font_name, start_size=11):
    size = start_size
    while size > 5:
        c.setFont(font_name, size)
        if c.stringWidth(text, font_name, size) <= max_width:
            return size
        size -= 0.5
    return size


def draw_cut_lines(c, x, y, w, h):
    """Linha tracejada nas bordas da etiqueta."""
    c.saveState()
    c.setStrokeColor(colors.Color(0.6, 0.6, 0.6))
    c.setLineWidth(0.3)
    c.setDash(3, 3)
    c.rect(x, y, w, h)
    c.restoreState()


def generate_pdf(entries):
    pdfmetrics.registerFont(TTFont('NotoSerifBold', FONT_PATH))
    pdfmetrics.registerFont(TTFont('NotoSans', FONT_SANS))

    c = canvas.Canvas(OUTPUT_PDF, pagesize=A4)
    total = len(entries)

    for idx, (cim, name, label_prefix) in enumerate(entries):
        col = idx % COLS
        row = (idx // COLS) % ROWS_PER_PAGE

        # Nova página
        if idx > 0 and col == 0 and row == 0:
            c.showPage()

        # Posição da célula (y cresce de baixo para cima no reportlab)
        x = MARGIN + col * CELL_W
        y = PAGE_H - MARGIN - (row + 1) * CELL_H

        # Linha de corte
        draw_cut_lines(c, x, y, CELL_W, CELL_H)

        # QR code centralizado horizontalmente, com margem superior
        qr_x = x + (CELL_W - QR_SIZE) / 2
        qr_y = y + CELL_H - QR_SIZE - 4 * mm

        qr_buf = make_qr_image(cim, QR_SIZE)
        c.drawImage(ImageReader(qr_buf), qr_x, qr_y, width=QR_SIZE, height=QR_SIZE)

        # CIM ou CPF abaixo do QR
        cim_y = qr_y - 5 * mm
        cim_text = f'{label_prefix}: {cim}'
        sz = fit_font_size(c, cim_text, CELL_W - 4 * mm, 'NotoSans', start_size=8)
        c.setFont('NotoSans', sz)
        c.setFillColor(colors.black)
        cim_w = c.stringWidth(cim_text, 'NotoSans', sz)
        c.drawString(x + (CELL_W - cim_w) / 2, cim_y, cim_text)

        # Nome abaixo do CIM
        name_y = cim_y - 4 * mm
        upper_name = name.upper()
        sz = fit_font_size(c, upper_name, CELL_W - 4 * mm, 'NotoSerifBold', start_size=9)
        c.setFont('NotoSerifBold', sz)
        name_w = c.stringWidth(upper_name, 'NotoSerifBold', sz)
        c.drawString(x + (CELL_W - name_w) / 2, name_y, upper_name)

        if (idx + 1) % 10 == 0 or (idx + 1) == total:
            print(f'  {idx+1}/{total}', end='\r')

    c.save()
    print(f'\nPDF gerado: {OUTPUT_PDF}')
    print(f'Total de etiquetas: {total} | Páginas: ~{((total // COLS) + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE}')


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    entries = []
    seen_names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[col['Nome']]).strip()
        if name.lower() in seen_names:
            continue

        cim = sanitize_cim(str(row[col['CIM']]))
        cpf = sanitize_cim(str(row[col['CPF']]))

        # CIM inválido: vazio ou mais de 7 dígitos (ex: CPF digitado no campo CIM)
        if cim and len(cim) <= 7:
            qr_data = cim
            label_prefix = 'CIM'
        elif cpf:
            qr_data = cpf
            label_prefix = 'CPF'
        else:
            continue

        seen_names.add(name.lower())
        entries.append((qr_data, name, label_prefix))

    # Ordem alfabética pelo nome
    entries.sort(key=lambda e: e[1].upper())

    print(f'{len(entries)} etiquetas para gerar (ordem alfabética)...')
    generate_pdf(entries)


if __name__ == '__main__':
    main()
