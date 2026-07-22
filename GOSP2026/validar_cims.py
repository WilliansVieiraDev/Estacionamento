#!/usr/bin/env python3
"""
Valida se os CIMs enviados via WhatsApp batem com os Codigos em appsettings.json.
"""
import json
import re
import openpyxl

APPSETTINGS = '/home/wvieira/projetos/GOSP/appsettings.json'
EXCEL_PATH  = '/home/wvieira/projetos/GOSP/produtos-vendidos-atualizado.xlsx'


def sanitize(v):
    return re.sub(r'\D', '', str(v))


# --- 1. Carrega Convidados do appsettings.json ---
with open(APPSETTINGS) as f:
    cfg = json.load(f)

convidados = cfg.get('Convidados', [])
# Mapeia CIM sanitizado → nome (appsettings)
app_map = {}
for entry in convidados:
    codigo_raw = entry.get('Codigo', '')
    codigo_san = sanitize(codigo_raw)
    if codigo_san:
        app_map[codigo_san] = entry.get('Nome', '')

print(f'appsettings.json: {len(convidados)} entradas, {len(app_map)} CIMs únicos sanitizados')

# --- 2. Carrega planilha e determina o que seria enviado ---
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb.active
headers = [cell.value for cell in ws[1]]
col = {h: i for i, h in enumerate(headers)}

sent_map  = {}   # CIM sanitizado → nome (planilha)
no_cim    = []   # Entradas sem CIM válido
seen_names = set()

for row in ws.iter_rows(min_row=2, values_only=True):
    name = str(row[col['Nome']]).strip()
    if name.lower() in seen_names:
        continue
    seen_names.add(name.lower())

    cim = sanitize(str(row[col['CIM']]))
    cpf = sanitize(str(row[col['CPF']]))

    if cim and len(cim) <= 7:
        sent_map[cim] = name
    elif cpf:
        # CPF → sem CIM, não seria enviado via WhatsApp
        no_cim.append((name, f'CPF={cpf}'))
    else:
        no_cim.append((name, 'sem CIM e sem CPF'))

print(f'Planilha: {len(sent_map)} com CIM válido, {len(no_cim)} sem CIM (CPF ou vazio)')

# --- 3. Comparação ---
app_cims  = set(app_map.keys())
sent_cims = set(sent_map.keys())

# CIMs enviados (planilha) que ESTÃO no appsettings
match     = sent_cims & app_cims
# CIMs enviados que NÃO estão no appsettings
not_in_app = sent_cims - app_cims
# CIMs do appsettings que NÃO foram enviados (não estão na planilha)
not_sent  = app_cims - sent_cims

print('\n' + '='*70)
print(f'  CORRESPONDÊNCIAS (enviados E presentes no appsettings): {len(match)}')
print(f'  Enviados mas AUSENTES do appsettings:                   {len(not_in_app)}')
print(f'  No appsettings mas NÃO na planilha (não enviados):      {len(not_sent)}')
print('='*70)

if not_in_app:
    print('\n--- Enviados via WhatsApp mas NÃO estão no appsettings.json ---')
    for cim in sorted(not_in_app, key=lambda c: sent_map[c]):
        print(f'  CIM={cim:<10} | {sent_map[cim]}')

if not_sent:
    print('\n--- No appsettings.json mas NÃO na planilha (não foram enviados) ---')
    for cim in sorted(not_sent):
        # Verifica se CIM é CPF (>7 dígitos) – já seriam descartados
        if len(cim) > 7:
            print(f'  CIM={cim:<15} | {app_map[cim]}  [CPF – sem envio WhatsApp]')
        else:
            print(f'  CIM={cim:<15} | {app_map[cim]}')

print('\n--- Sem CIM válido na planilha (não enviados) ---')
for name, motivo in no_cim:
    print(f'  {name:<50} {motivo}')
