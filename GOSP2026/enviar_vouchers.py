#!/usr/bin/env python3
import openpyxl
import smtplib
import re
import time
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

# ── Configurações ──────────────────────────────────────────────────────────────
EXCEL_PATH   = '/home/wvieira/projetos/GOSP/produtos-vendidos-final.xlsx'
VOUCHER_BASE = 'https://estacionamento-beta-rosy.vercel.app/img/{cim}.jpg'

SMTP_HOST = 'smtp.gmail.com'
SMTP_PORT = 587
SMTP_USER = 'francis.bassi@gosp.org.br'
SMTP_PASS = 'rmkf xskp lwiv obge'

# Antes de disparar para todos, envie TEST_MODE=True para validar
TEST_MODE      = False
TEST_CIM       = '17668'  # CIM do destinatário usado no teste
TEST_EMAILS    = ['dev.willians@gmail.com', 'lpfernandes_@hotmail.com']

DELAY_SECONDS  = 2   # pausa entre envios para não cair em spam


def sanitize_cim(cim: str) -> str:
    return re.sub(r'\D', '', cim)


def build_email(name: str, voucher_url: str) -> tuple[str, str]:
    """Retorna (assunto, html)."""
    first = name.strip().split()[0].capitalize()

    subject = 'Seu Voucher de Acesso — Congresso GOSP 2026'

    html = f"""\
<html><body style="font-family:Arial,sans-serif;font-size:15px;color:#222;line-height:1.6">
<p>Meu Irmão,</p>

<p>Segue o seu voucher com o QR Code individual de acesso ao <strong>Congresso GOSP 2026</strong>.</p>

<div style="text-align:center;margin:28px 0;">
  <a href="{voucher_url}" style="
    display:inline-block;
    background:linear-gradient(180deg,#1a2340 0%,#0d1526 100%);
    border:2px solid #b8952a;
    border-radius:6px;
    padding:14px 30px;
    text-decoration:none;
    font-family:Georgia,serif;
    letter-spacing:2px;
    line-height:1.3;
    text-align:center;
  ">
    <span style="display:block;font-size:22px;font-weight:bold;color:#d4a017;letter-spacing:4px;">VOUCHER</span>
    <span style="display:block;font-size:11px;color:#c9a84c;letter-spacing:3px;margin-top:2px;">CONGRESSO GOSP 2026</span>
  </a>
</div>

<p>Salve a imagem no celular e apresente o QR Code na entrada para a confirmação da inscrição e a retirada do kit oficial do Congresso.</p>

<p>O QR Code é <strong>pessoal e intransferível</strong>.</p>

<hr style="border:none;border-top:1px solid #ddd;margin:20px 0">

<p>📅 <strong>Credenciamento e retirada do kit</strong></p>
<p>
  Sexta-feira, 24/07<br>A partir das 18h<br><br>
  Sábado, 25/07<br>A partir das 8h
</p>

<p>📍 <strong>Grande Oriente de São Paulo</strong><br>
Rua São Joaquim, 457 — Liberdade<br>
São Paulo — SP<br>
CEP 01508-001</p>

<hr style="border:none;border-top:1px solid #ddd;margin:20px 0">

<p>Para facilitar a leitura, deixe o voucher aberto e aumente o brilho da tela do celular antes de chegar ao atendimento.</p>

<p><em>Congresso GOSP 2026<br>
Quando a Maçonaria fala ao mundo, o mundo escuta.</em></p>
</body></html>"""

    return subject, html


def send(smtp, to_addr: str, name: str, voucher_url: str, dry_run_to: list[str] = None):
    subject, html = build_email(name, voucher_url)

    msg = MIMEMultipart('alternative')
    msg['Subject'] = subject
    msg['From']    = f'Congresso GOSP 2026 <{SMTP_USER}>'
    msg['To']      = to_addr if not dry_run_to else ', '.join(dry_run_to)

    msg.attach(MIMEText(html, 'html', 'utf-8'))

    recipients = dry_run_to if dry_run_to else [to_addr]
    smtp.sendmail(SMTP_USER, recipients, msg.as_bytes())


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cim   = sanitize_cim(str(row[col['CIM']]))
        name  = str(row[col['Nome']])
        email = str(row[col['Email']]).strip()
        if not cim:
            print(f'  [SKIP] {name} — sem CIM')
            continue
        if not email or '@' not in email:
            print(f'  [SKIP] {name} — email inválido: {email}')
            continue
        rows.append((cim, name, email))

    print(f'\n{len(rows)} destinatários encontrados.')

    if TEST_MODE:
        print(f'\n*** MODO TESTE — enviando apenas para {TEST_EMAILS} ***\n')

    print(f'Conectando a {SMTP_HOST}:{SMTP_PORT}...')
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as smtp:
        smtp.ehlo()
        smtp.starttls()
        smtp.login(SMTP_USER, SMTP_PASS)
        print('Autenticado.\n')

        if TEST_MODE:
            # Busca o destinatário de teste pelo CIM ou usa o primeiro da lista
            test_row = next((r for r in rows if r[0] == TEST_CIM), rows[0])
            cim, name, email = test_row
            voucher_url = VOUCHER_BASE.format(cim=cim)
            send(smtp, email, name, voucher_url, dry_run_to=TEST_EMAILS)
            print(f'[TESTE] → {TEST_EMAILS} | CIM={cim} | {name} | voucher={voucher_url}')
            print('\nVerifique os emails de teste. Quando aprovado, defina TEST_MODE=False e rode novamente.')
            return

        sent = 0
        for cim, name, email in rows:
            voucher_url = VOUCHER_BASE.format(cim=cim)
            try:
                send(smtp, email, name, voucher_url)
                print(f'[{sent+1}/{len(rows)}] ✓ {name} <{email}>')
                sent += 1
                time.sleep(DELAY_SECONDS)
            except Exception as e:
                print(f'[ERRO] {name} <{email}>: {e}')

    print(f'\n{sent} emails enviados.')


if __name__ == '__main__':
    main()
